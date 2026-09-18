"""Structure-aware OOXML readers used after the package has been validated.

These readers never resolve an external relationship or recalculate a formula.
Spreadsheet values are labelled with their stored format rather than promising
pixel-equivalent Excel display rendering.
"""

import posixpath
import re
import zipfile
from datetime import date, datetime, time
from itertools import zip_longest
from pathlib import Path

from defusedxml import ElementTree as XML

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
DOCX_NOTICE = "[이미지와 차트의 시각 정보 및 원본 페이지 배치는 포함하지 않습니다. 텍스트 상자의 문자열은 추출합니다.]"
XLSX_NOTICE = "[추출 범위: 시트와 셀 좌표별 저장 값, 수식 및 저장된 계산 결과. 수식을 재계산하지 않습니다. 숫자와 날짜는 저장 값을 기준으로 표시하며 표시 형식을 별도로 제공합니다. 조건부 서식, 차트와 이미지의 시각 정보는 포함하지 않습니다.]"


def _xml(archive, name):
    root = XML.fromstring(archive.read(name), forbid_dtd=True)
    pending, count = [(root, 0)], 0
    while pending:
        node, depth = pending.pop()
        count += 1
        if count > 200000 or depth > 128:
            raise ValueError("OFFICE_SIZE_LIMIT")
        pending.extend((child, depth + 1) for child in node)
    return root


def _children(node):
    for child in node:
        if child.tag == MC + "AlternateContent":
            branch = child.find(MC + "Choice")
            if branch is None:
                branch = child.find(MC + "Fallback")
            if branch is not None:
                yield from _children(branch)
        else:
            yield child


def _inline(node):
    if node.tag in {W + "t", M + "t"}:
        return node.text or ""
    if node.tag == W + "tab":
        return "\t"
    if node.tag in {W + "br", W + "cr"}:
        return "\n"
    if node.tag in {W + "footnoteReference", W + "endnoteReference"}:
        label = "각주" if node.tag == W + "footnoteReference" else "미주"
        return f"[{label} {node.get(W + 'id', '')}]"
    # A text box can contain paragraphs or a table inside a drawing in a run.
    # Consume its block tree here exactly once instead of iterating all w:p.
    if node.tag == W + "txbxContent":
        content = "\n".join(_blocks(node))
        return f"\n[텍스트 상자]\n{content}\n" if content else ""
    if node.tag in {W + "del", W + "moveFrom"}:
        return ""
    return "".join(_inline(child) for child in _children(node))


def _table_children(node, tag):
    for child in _children(node):
        if child.tag == W + tag:
            yield child
        elif child.tag not in {W + "tbl", W + "tr", W + "tc"}:
            yield from _table_children(child, tag)


def _blocks(node):
    for child in _children(node):
        if child.tag == W + "p":
            text = _inline(child).strip()
            if text:
                yield text
        elif child.tag == W + "tbl":
            for row_index, row in enumerate(_table_children(child, "tr"), 1):
                cells, column, has_text = [], 1, False
                before = row.find(f"{W}trPr/{W}gridBefore")
                if before is not None:
                    try:
                        column += max(0, int(before.get(W + "val", "0")))
                    except ValueError:
                        pass
                for cell in _table_children(row, "tc"):
                    text = "\n".join(_blocks(cell))
                    has_text = has_text or bool(text.strip())
                    span = cell.find(f"{W}tcPr/{W}gridSpan")
                    try:
                        width = max(1, int(span.get(W + "val", "1"))) if span is not None else 1
                    except ValueError:
                        width = 1
                    merge = cell.find(f"{W}tcPr/{W}vMerge")
                    label = f"열 {column}" if width == 1 else f"열 {column}~{column + width - 1}"
                    if merge is not None and merge.get(W + "val") != "restart":
                        label += " (위 셀과 병합)"
                    cells.append(f"[{label}] {text}")
                    column += width
                if has_text:
                    yield f"[표 행 {row_index}] " + " | ".join(cells)
        elif child.tag not in {W + "del", W + "moveFrom"}:
            yield from _blocks(child)


def _docx_content(archive):
    document = _xml(archive, "word/document.xml")
    body = document.find(W + "body")
    if body is not None:
        yield from _blocks(body)
    relationships, note_parts = {}, {}
    rel_path = "word/_rels/document.xml.rels"
    if rel_path in archive.namelist():
        for rel in _xml(archive, rel_path):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target", "")
            normalized = posixpath.normpath(posixpath.join("word", target))
            if target.startswith("/"):
                normalized = posixpath.normpath(target.lstrip("/"))
            if normalized.startswith("../") or normalized not in archive.namelist():
                continue
            relationships[rel.get("Id")] = normalized
            kind = rel.get("Type", "").rsplit("/", 1)[-1]
            if kind in {"footnotes", "endnotes"}:
                note_parts[kind] = normalized
    seen = set()
    for node in document.iter():
        if node.tag not in {W + "headerReference", W + "footerReference"}:
            continue
        part = relationships.get(node.get(R + "id"))
        if not part or part in seen:
            continue
        seen.add(part)
        label = "머리말" if node.tag == W + "headerReference" else "꼬리말"
        for text in _blocks(_xml(archive, part)):
            yield f"[{label}] {text}"
    for part, tag, label in ((note_parts.get("footnotes", "word/footnotes.xml"), "footnote", "각주"),
                             (note_parts.get("endnotes", "word/endnotes.xml"), "endnote", "미주")):
        if part not in archive.namelist():
            continue
        for note in _xml(archive, part):
            if note.tag != W + tag or note.get(W + "type", "normal") != "normal":
                continue
            note_id = note.get(W + "id", "")
            if note_id.startswith("-"):
                continue
            for text in _blocks(note):
                yield f"[{label} {note_id}] {text}"


def docx_parts(path: Path):
    """Yield body-order blocks and linked supplementary document text."""
    with zipfile.ZipFile(path) as archive:
        visual = any(
            re.search(rb"<(?:[\w.-]+:)?(?:drawing|pict|object)(?:\s|>)", archive.read(name))
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
        first = True
        for text in _docx_content(archive):
            if text.strip():
                if first and visual:
                    yield DOCX_NOTICE
                first = False
                yield text


def _value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    # Preserve line breaks and delimiters unambiguously in row-oriented output.
    return str(value).replace("\\", "\\\\").replace("\r", "\\r").replace("\n", "\\n").replace(" | ", " \\| ")


def _display(value, number_format):
    """Render only simple, unambiguous numeric formats; retain all originals."""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    match = re.fullmatch(r'([₩$€£]?)([#,0]+(?:\.0+)?)(%?)(?:"([^";]*)"|((?:\\.)+))?', number_format)
    if not match:
        return None
    prefix, digits, percent, suffix, escaped_suffix = match.groups()
    if digits.partition(".")[0].count("0") > 1:
        return None  # Keep the original format; zero-padding is not rendered here.
    if escaped_suffix:
        suffix = re.sub(r"\\(.)", r"\1", escaped_suffix)
    decimals = len(digits.partition(".")[2])
    number = value * 100 if percent else value
    formatted = format(number, f"{',' if ',' in digits else ''}.{decimals}f")
    return prefix + formatted + percent + (suffix or "")


def _cell_text(cell, cached):
    value = cell.value
    if cell.data_type == "f":
        formula = value if isinstance(value, str) else getattr(value, "text", None)
        formula = _value(formula) if formula is not None else "[특수 수식 원문 없음]"
        result = cached.value if cached is not None else None
        text = f"{formula}; 수식의 저장된 결과: {_value(result) if result is not None else '[없음: 재계산 필요]'}"
        display_value = result
    else:
        text = _value(value)
        display_value = value
    number_format = cell.number_format
    if number_format and number_format != "General":
        text += f"; 표시 형식: {number_format}"
        rendered = _display(display_value, number_format)
        if rendered is not None:
            text += f"; 표시 값: {rendered}"
    return f"{cell.coordinate}: {text}"


def xlsx_parts(path: Path):
    """Keep sheet/row/cell identities and pair formulas with saved results."""
    from openpyxl import load_workbook

    with path.open("rb") as formula_handle, path.open("rb") as value_handle:
        formulas = load_workbook(formula_handle, read_only=True, data_only=False, keep_links=False)
        values = None
        try:
            values = load_workbook(value_handle, read_only=True, data_only=True, keep_links=False)
            cells, first = 0, True
            for sheet in formulas:
                if ((sheet.max_row and sheet.max_row > 100000)
                        or (sheet.max_column and sheet.max_column > 2000)):
                    raise ValueError("OFFICE_SIZE_LIMIT")
                cached_sheet = values[sheet.title]
                sheet.reset_dimensions()
                cached_sheet.reset_dimensions()
                sheet_started = False
                for index, rows in enumerate(zip_longest(sheet.iter_rows(), cached_sheet.iter_rows(), fillvalue=()), 1):
                    row, cached_row = rows
                    cells += max(len(row), len(cached_row))
                    if cells > 200000 or index > 100000 or len(row) > 2000:
                        raise ValueError("OFFICE_SIZE_LIMIT")
                    entries = [_cell_text(cell, cached_row[offset] if offset < len(cached_row) else None)
                               for offset, cell in enumerate(row) if cell.value is not None and str(cell.value).strip()]
                    if not entries:
                        continue
                    if first:
                        yield XLSX_NOTICE
                        first = False
                    if not sheet_started:
                        yield f"[시트: {sheet.title}]"
                        sheet_started = True
                    yield f"[행 {index}] " + " | ".join(entries)
        finally:
            formulas.close()
            if values is not None:
                values.close()
