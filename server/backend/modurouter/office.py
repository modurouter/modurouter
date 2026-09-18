"""Local OOXML text extraction. Never execute formulas, macros or external links."""

import zipfile
from pathlib import Path, PurePosixPath
from xml.etree.ElementTree import ParseError

from defusedxml import ElementTree as XML
from defusedxml.common import DefusedXmlException

from .document_formats import DOCX, MAX_TEXT, OFFICE_PARTS, PPTX
from .office_fidelity import docx_parts, xlsx_parts

MAX_EXPANDED = 64 * 1024**2
MAX_MEMBER = 16 * 1024**2
ERRORS = {"OFFICE_INVALID", "OFFICE_ENCRYPTED", "OFFICE_SIZE_LIMIT"}


def validate_office(path: Path, mime: str, *, inspect_xml: bool = False):
    with path.open("rb") as handle:
        if handle.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ValueError("OFFICE_ENCRYPTED")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if len(entries) > 4096 or sum(e.file_size for e in entries) > MAX_EXPANDED or any(e.file_size > MAX_MEMBER for e in entries):
                raise ValueError("OFFICE_SIZE_LIMIT")
            if len(names) != len(entries) or any(PurePosixPath(n).is_absolute() or ".." in PurePosixPath(n).parts for n in names):
                raise ValueError("OFFICE_INVALID")
            if any(e.flag_bits & 1 for e in entries):
                raise ValueError("OFFICE_ENCRYPTED")
            part, content_type = OFFICE_PARTS[mime]
            if part not in names or "[Content_Types].xml" not in names:
                raise ValueError("FILE_TYPE_MISMATCH")
            types = XML.fromstring(archive.read("[Content_Types].xml"), forbid_dtd=True)
            if not any(e.get("PartName") == "/" + part and e.get("ContentType") == content_type for e in types):
                raise ValueError("FILE_TYPE_MISMATCH")
            if inspect_xml:
                for name in names:
                    if name.endswith((".xml", ".rels")):
                        XML.fromstring(archive.read(name), forbid_dtd=True)
    except (OSError, zipfile.BadZipFile, KeyError, ParseError, DefusedXmlException, NotImplementedError) as exc:
        raise ValueError("OFFICE_INVALID") from exc


def _pptx(path):
    from pptx import Presentation

    def shapes_text(shapes):
        for shape in shapes:
            if shape.has_text_frame:
                yield shape.text_frame.text
            if shape.has_table:
                for row in shape.table.rows:
                    values = [cell.text for cell in row.cells]
                    if any(value.strip() for value in values):
                        yield " | ".join(values)
            if hasattr(shape, "shapes"):
                yield from shapes_text(shape.shapes)

    presentation = Presentation(path)
    for index, slide in enumerate(presentation.slides, 1):
        announced = False
        for text in shapes_text(slide.shapes):
            if text.strip():
                if not announced:
                    yield f"[슬라이드 {index}]"
                    announced = True
                yield text
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                if not announced:
                    yield f"[슬라이드 {index}]"
                yield "[발표자 노트]\n" + notes


def office_text(path: Path, mime: str) -> str:
    validate_office(path, mime, inspect_xml=True)
    reader = docx_parts if mime == DOCX else _pptx if mime == PPTX else xlsx_parts
    parts, length = [], 0
    iterator = reader(path)
    try:
        for part in iterator:
            if not part.strip():
                continue
            length += len(part) + bool(parts)
            parts.append(part)
            if length > MAX_TEXT:
                break
    finally:
        iterator.close()
    return "\n".join(parts)
