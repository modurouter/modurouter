import zipfile

import pytest
from docx import Document
from modurouter.document_formats import DOCX, MAX_TEXT, PPTX, XLSX
from modurouter.extract import extract
from modurouter.office import validate_office
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def rewrite_archive(path, changes):
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    data.update(changes)
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in data.items():
            archive.writestr(name, content)


def test_docx_paragraphs_nested_tables_headers_and_footer(tmp_path):
    path = tmp_path / 'storage-key'
    document = Document()
    document.sections[0].header.paragraphs[0].text = '문서 머리말'
    document.add_paragraph('첫 번째 본문')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '첫 셀'
    table.cell(0, 1).text = '둘째 셀'
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = '중첩 셀'
    document.add_paragraph('마지막 본문')
    document.sections[0].footer.paragraphs[0].text = '문서 꼬리말'
    document.save(path)
    result = extract(path, DOCX)
    text = result['text']
    for part in ['첫 번째 본문', '첫 셀', '둘째 셀', '중첩 셀', '마지막 본문', '문서 머리말', '문서 꼬리말']:
        assert text.count(part) == 1
    assert text.index('첫 번째 본문') < text.index('첫 셀') < text.index('마지막 본문')
    assert not result['truncated']


def test_xlsx_multiple_sheets_formulas_and_incorrect_dimensions(tmp_path):
    path = tmp_path / 'storage-key'
    book = Workbook()
    book.active.title = '매출'
    book.active.append(['품목', '수량', '금액'])
    book.active.append(['책', 3, '=B2*100'])
    book.create_sheet('메모')['B3'] = '다음 분기'
    book.save(path)
    with zipfile.ZipFile(path) as archive:
        sheet = archive.read('xl/worksheets/sheet1.xml').replace(b'ref="A1:C2"', b'ref="A1:A1"')
    rewrite_archive(path, {'xl/worksheets/sheet1.xml': sheet})
    result = extract(path, XLSX)
    assert all(part in result['text'] for part in ['매출', 'A2: 책', 'B2: 3', 'C2: =B2*100', '메모', 'B3: 다음 분기'])
    assert not result['truncated']


def test_pptx_text_tables_groups_and_notes(tmp_path):
    path = tmp_path / 'storage-key'
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = '제목과 본문'
    table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(2), Inches(1)).table
    table.cell(0, 0).text = '표 왼쪽'
    table.cell(0, 1).text = '표 오른쪽'
    slide.shapes.add_group_shape().shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = '그룹 텍스트'
    slide.notes_slide.notes_text_frame.text = '발표자 설명'
    presentation.save(path)
    result = extract(path, PPTX)
    assert all(part in result['text'] for part in ['슬라이드 1', '제목과 본문', '표 왼쪽', '표 오른쪽', '그룹 텍스트', '발표자 설명'])
    assert not result['truncated']


@pytest.mark.parametrize('mime', [DOCX, XLSX, PPTX])
def test_empty_office_is_not_marker_only_success(tmp_path, mime):
    path = tmp_path / 'storage-key'
    document = Document() if mime == DOCX else Workbook() if mime == XLSX else Presentation()
    if mime == PPTX:
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_table(1, 2, 0, 0, Inches(2), Inches(1))
    document.save(path)
    with pytest.raises(ValueError, match='NO_EXTRACTABLE_TEXT'):
        extract(path, mime)


@pytest.mark.parametrize('length,truncated', [(MAX_TEXT, False), (MAX_TEXT + 1, True)])
def test_docx_truncation_is_exact(tmp_path, length, truncated):
    path = tmp_path / 'storage-key'
    document = Document()
    document.add_paragraph('가' * length)
    document.save(path)
    assert extract(path, DOCX) == {'text': '가' * MAX_TEXT, 'truncated': truncated}


def test_docx_does_not_stop_at_exact_boundary_before_next_paragraph(tmp_path):
    path = tmp_path / 'storage-key'
    document = Document()
    document.add_paragraph('가' * MAX_TEXT)
    document.add_paragraph('마지막')
    document.save(path)
    assert extract(path, DOCX)['truncated']


@pytest.mark.parametrize('payload,error', [(b'not a ZIP', 'OFFICE_INVALID'),
    (b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest', 'OFFICE_ENCRYPTED')])
def test_invalid_office_containers(tmp_path, payload, error):
    path = tmp_path / 'storage-key'
    path.write_bytes(payload)
    with pytest.raises(ValueError, match=error):
        extract(path, DOCX)


def test_office_mime_mismatch(tmp_path):
    path = tmp_path / 'storage-key'
    Document().save(path)
    with pytest.raises(ValueError, match='FILE_TYPE_MISMATCH'):
        extract(path, XLSX)


@pytest.mark.parametrize('extra', [
    {'word/evil.xml': '<!DOCTYPE x [<!ENTITY x "unsafe">]><x>&x;</x>'},
    {'word/evil.xml': '<broken>'},
    {'../evil': 'text'},
    {'/evil': 'text'},
])
def test_office_malformed_dtd_and_unsafe_members(tmp_path, extra):
    path = tmp_path / 'storage-key'
    Document().save(path)
    rewrite_archive(path, extra)
    with pytest.raises(ValueError, match='OFFICE_INVALID'):
        extract(path, DOCX)


def test_office_zip_bomb_expansion_limit(tmp_path):
    path = tmp_path / 'storage-key'
    Document().save(path)
    rewrite_archive(path, {'word/bomb.xml': b' ' * (16 * 1024**2 + 1)})
    assert path.stat().st_size < 100000
    with pytest.raises(ValueError, match='OFFICE_SIZE_LIMIT'):
        validate_office(path, DOCX)


def test_office_duplicate_member_is_rejected(tmp_path):
    path = tmp_path / 'storage-key'
    Document().save(path)
    with pytest.warns(UserWarning, match='Duplicate name'), zipfile.ZipFile(path, 'a') as archive:
        archive.writestr('word/document.xml', '<x/>')
    with pytest.raises(ValueError, match='OFFICE_INVALID'):
        extract(path, DOCX)
