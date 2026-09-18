"""Create original, redistributable issue #6 extraction fixtures.

Run from server: uv run python .runtime/issue6-fixtures/generate.py
All document content is synthetic and contains no user data.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
CASES = []


def case(filename, mime, expected, *, ordered=(), forbidden=()):
    CASES.append(dict(filename=filename, mime=mime, expected=list(expected),
                      ordered=list(ordered), forbidden=list(forbidden)))


def rewrite_zip(path, updates):
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries.update(updates)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


def docx():
    document = Document()
    document.add_heading('모두라우터 문서 검증 보고서', 0)
    document.sections[0].header.paragraphs[0].text = '머리말: 검증부서 청록팀'
    document.add_paragraph('문서처음: 시작코드는 가람101입니다.')
    document.add_paragraph('일정은 2026년 9월 18일이며 단위는 원입니다.')
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = '사업명'
    table.rows[0].cells[1].text = '예산(원)'
    table.rows[0].cells[2].text = '담당자'
    for values in [('달빛사업', '1250000', '김하늘'), ('별빛사업', '2500000', '이바다')]:
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = value
    document.add_paragraph('문서중간: 중간코드는 나래202입니다.')
    paragraph = document.add_paragraph('본문의 각주 참조')
    reference = OxmlElement('w:footnoteReference')
    reference.set(qn('w:id'), '1')
    paragraph.add_run()._r.append(reference)
    textbox = parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape id="Issue6TextBox" style="width:240pt;height:30pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>글상자: 별도코드는 다솜303입니다.</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>''')
    document.add_paragraph()._p.append(textbox)
    document.add_page_break()
    document.add_paragraph('문서끝: 종료코드는 라온404입니다.')
    document.sections[0].footer.paragraphs[0].text = '꼬리말: 승인번호 마루505'
    path = ROOT / 'korean-report.docx'
    document.save(path)
    with ZipFile(path) as archive:
        types = archive.read('[Content_Types].xml').decode()
        rels = archive.read('word/_rels/document.xml.rels').decode()
    types = types.replace('</Types>', '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
    rels = rels.replace('</Relationships>', '<Relationship Id="rIdIssue6Footnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
    footnotes = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote><w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>각주: 계약확인 코드는 바람606입니다.</w:t></w:r></w:p></w:footnote></w:footnotes>'''
    rewrite_zip(path, {'[Content_Types].xml': types, 'word/_rels/document.xml.rels': rels, 'word/footnotes.xml': footnotes})
    case(path.name, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         ['가람101', '나래202', '라온404', '달빛사업', '1250000', '김하늘', '별빛사업', '이바다', '다솜303', '바람606', '마루505', '청록팀'],
         ordered=['가람101', '달빛사업', '별빛사업', '나래202', '라온404'])


def xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '분기예산'
    sheet.append(['항목', '수량(개)', '단가(원)', '합계(원)', '기준일'])
    sheet.append(['첫행 가람101', 3, 1200, '=B2*C2', datetime(2026, 9, 18)])
    sheet.append(['중간행 나래202', 2, 2500, '=B3*C3', datetime(2026, 9, 19)])
    sheet.append(['마지막행 라온404', 1, 3000, '=B4*C4', datetime(2026, 9, 20)])
    sheet['E2'].number_format = 'yyyy-mm-dd'
    sheet['E3'].number_format = 'yyyy-mm-dd'
    sheet['E4'].number_format = 'yyyy-mm-dd'
    for row in (2, 3, 4):
        sheet[f'C{row}'].number_format = '#,##0"원"'
        sheet[f'D{row}'].number_format = '#,##0"원"'
    second = workbook.create_sheet('배송현황')
    second.append(['배송처', '예정일', '비고'])
    second.append(['제주센터', datetime(2026, 10, 2), '다른시트 코드 사과707'])
    second['B2'].number_format = 'yyyy-mm-dd'
    second.append(['서울센터', datetime(2026, 10, 3), '두 줄 메모\n최종확인 코드 아람808'])
    second['B3'].number_format = 'yyyy-mm-dd'
    hidden = workbook.create_sheet('비공개표시시트')
    hidden.append(['숨김시트 안내', '숨김내용 자람909'])
    hidden.sheet_state = 'hidden'
    path = ROOT / 'korean-budget.xlsx'
    workbook.save(path)
    with ZipFile(path) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').decode()
    for formula, cached in [('B2*C2', '3600'), ('B3*C3', '5000'), ('B4*C4', '3000')]:
        xml = xml.replace(f'<f>{formula}</f><v></v>', f'<f>{formula}</f><v>{cached}</v>')
    rewrite_zip(path, {'xl/worksheets/sheet1.xml': xml})
    case(path.name, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
         ['분기예산', '배송현황', '가람101', '나래202', '라온404', '사과707', '아람808', '2026-09-18', '수량(개)', '단가(원)', '3600', 'B2*C2', '자람909'],
         ordered=['가람101', '나래202', '라온404', '사과707', '아람808'])


def pptx():
    presentation = Presentation()
    entries = [
        ('첫 번째 슬라이드', '발표처음: 가람101', '발표자메모 첫째: 회의실 청록호'),
        ('두 번째 슬라이드', '발표중간: 나래202', '발표자메모 둘째: 승인코드 바람606'),
        ('세 번째 슬라이드', '발표끝: 라온404', '발표자메모 셋째: 종료시각 18시'),
    ]
    for index, (title, body, notes) in enumerate(entries):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
        if index == 1:
            table = slide.shapes.add_table(3, 2, Inches(1), Inches(4), Inches(6), Inches(1.5)).table
            for row, values in enumerate([('항목', '수치'), ('전환율', '27%'), ('만족도', '92점')]):
                for column, value in enumerate(values):
                    table.cell(row, column).text = value
    path = ROOT / 'korean-slides.pptx'
    presentation.save(path)
    case(path.name, 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
         ['가람101', '나래202', '라온404', '전환율', '27%', '만족도', '92점', '청록호', '바람606', '18시'],
         ordered=['가람101', '나래202', '라온404'])


def text_formats():
    (ROOT / 'korean-notes.md').write_text('# 검증 문서\n\n처음: 가람101\n\n| 항목 | 값 |\n| --- | --- |\n| 표코드 | 나래202 |\n\n```python\nprint("인용된 코드")\n```\n\n끝: 라온404\n', encoding='utf-8')
    case('korean-notes.md', 'text/plain', ['가람101', '나래202', '라온404', '인용된 코드'], ordered=['가람101', '나래202', '라온404'])
    html = '<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>검증 HTML 문서</title><style>.hidden{content:"STYLE_MUST_NOT_LEAK"}</style><script>const secret="SCRIPT_MUST_NOT_LEAK";</script></head><body><h1>처음: 가람101</h1><p>중간 전 단락</p><table><tr><th>사업명</th><th>예산</th></tr><tr><td>나래202</td><td>5,000원</td></tr></table><p>끝: 라온404</p><img src="https://invalid.example/image.png" alt="그림 설명 바람606"></body></html>'
    (ROOT / 'korean-page.html').write_text(html, encoding='utf-8')
    case('korean-page.html', 'text/html', ['가람101', '나래202', '라온404', '5,000원'], ordered=['가람101', '나래202', '라온404'], forbidden=['SCRIPT_MUST_NOT_LEAK', 'STYLE_MUST_NOT_LEAK'])
    for suffix, delimiter in [('csv', ','), ('tsv', '\t')]:
        output = io.StringIO(newline='')
        writer = csv.writer(output, delimiter=delimiter)
        writer.writerows([['구분', '메모', '수치(원)'], ['처음 가람101', '쉼표, 포함', '1,000'], ['중간 나래202', '줄바꿈\n둘째 줄', '2500'], ['끝 라온404', '인용 "발표"', '3500']])
        (ROOT / f'korean-data.{suffix}').write_text(output.getvalue(), encoding='utf-8-sig', newline='')
        case(f'korean-data.{suffix}', 'text/plain', ['가람101', '나래202', '라온404', '쉼표', '둘째 줄'], ordered=['가람101', '나래202', '라온404'])
    (ROOT / 'korean-cp949.csv').write_bytes('항목,설명\r\n가람101,한글 인코딩 검증\r\n나래202,서울시 강남구\r\n라온404,마지막 데이터\r\n'.encode('cp949'))
    case('korean-cp949.csv', 'text/plain', ['가람101', '나래202', '라온404', '한글 인코딩 검증'])
    (ROOT / 'corrupt.docx').write_bytes(b'PK\x03\x04broken-container')
    (ROOT / 'renamed-text.docx').write_text('이 파일은 DOCX가 아니라 일반 텍스트입니다.', encoding='utf-8')
    (ROOT / 'renamed-office.xlsx').write_bytes((ROOT / 'korean-report.docx').read_bytes())


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    docx()
    xlsx()
    pptx()
    text_formats()
    (ROOT / 'manifest.json').write_text(json.dumps(CASES, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'fixtures': len(CASES), 'directory': str(ROOT)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
