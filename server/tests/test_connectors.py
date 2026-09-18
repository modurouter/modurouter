import asyncio
import io
import json

import pytest
import test_harness as harness_tests
from modurouter import files, harness, worker
from modurouter.errors import AppError
from modurouter.retrieval import requested_urls, should_search

world = harness_tests.world
provider = harness_tests.provider


@pytest.mark.parametrize("question,mode,attachment,expected", [
    ("오늘 서울 날씨 알려줘", None, False, True),
    ("최신 정책을 찾아줘", False, False, False),
    ("일반 질문", True, False, True),
    ("이 문서의 최근 매출 요약", None, True, False),
    ("현재 상태를 설명해줘", None, True, False),
    ("이 문서와 웹 검색 자료를 비교", None, True, True),
    ("https://example.com 문서를 요약", None, False, False),
    ("검색하지 말고 오늘 일기 써줘", None, False, False),
    ("이 문장 번역해줘", None, False, False),
])
def test_automatic_search_respects_mode_and_document_privacy(question, mode, attachment, expected):
    assert should_search(question, mode, has_attachments=attachment) is expected


def test_question_urls_are_deduplicated():
    assert requested_urls("https://example.com/a https://example.com/a. (https://example.com/b)") == [
        "https://example.com/a", "https://example.com/b"]


async def test_auto_search_reads_multiple_sources_with_text_only_model(world, provider, monkeypatch):
    monkeypatch.setattr(harness.settings, "tool_model_allowlist", "")
    read, searched = [], []

    async def search(query):
        searched.append(query)
        return [{"title": f"결과 {i}", "url": f"https://example.com/{i}", "text": "검색 발췌",
                 "scope": "search_snippet"} for i in range(5)]

    async def page(url):
        read.append(url)
        if url.endswith("/0"):
            raise AppError("PAGE_UNAVAILABLE", "페이지 실패")
        return {"title": url, "url": url, "text": "검증된 본문 " + url, "scope": "page"}

    monkeypatch.setattr(harness, "search_web", search)
    monkeypatch.setattr(harness, "read_url", page)
    client, cid = world
    response = await client.post(f"/v1/conversations/{cid}/runs", json={"message": "오늘 최신 뉴스 알려줘"})
    assert '"status": "completed"' in response.text
    assert searched == ["오늘 최신 뉴스 알려줘"]
    assert len(read) == 4 and len(provider.calls) == 1
    assert "일부 링크의 원문을 읽지 못했습니다" in response.text
    contents = "\n".join(m["content"] for m in provider.messages[-1])
    for i in (1, 2, 3):
        assert "검증된 본문 https://example.com/" + str(i) in contents


async def test_explicit_links_are_read_even_with_search_enabled(world, provider, monkeypatch):
    monkeypatch.setattr(harness.settings, "tool_model_allowlist", "")
    calls = []

    async def page(url):
        calls.append(url)
        return {"title": url, "url": url, "text": "문서 본문", "scope": "page"}

    async def search(query):
        calls.append("search")
        return [{"title": "결과", "url": "https://example.com/result", "text": "발췌", "scope": "search_snippet"}]

    monkeypatch.setattr(harness, "read_url", page)
    monkeypatch.setattr(harness, "search_web", search)
    client, cid = world
    response = await client.post(f"/v1/conversations/{cid}/runs", json={
        "message": "https://example.com/a.pdf https://example.com/b.docx 비교해줘", "search_enabled": True})
    assert '"status": "completed"' in response.text
    assert set(calls[:2]) == {"https://example.com/a.pdf", "https://example.com/b.docx"}
    assert calls[2:] == ["search", "https://example.com/result"]


async def test_search_disabled_never_calls_search_but_reads_urls(world, provider, monkeypatch):
    monkeypatch.setattr(harness.settings, "tool_model_allowlist", "")

    async def forbidden(query):
        pytest.fail("Disabled search must never leave the server")

    async def page(url):
        return {"title": "공식 문서", "url": url, "text": "공식 본문", "scope": "page"}

    monkeypatch.setattr(harness, "read_url", page)
    monkeypatch.setattr(harness, "search_web", forbidden)
    client, cid = world
    response = await client.post(f"/v1/conversations/{cid}/runs", json={
        "message": "https://example.com 최신 자료 비교", "search_enabled": False})
    assert '"status": "completed"' in response.text and len(provider.calls) == 1


@pytest.mark.parametrize("suffix", ["docx", "xlsx", "pptx", "md", "csv", "json", "xml"])
async def test_document_upload_worker_preview_and_text_only_followup(world, provider, database, monkeypatch, tmp_path, suffix):
    monkeypatch.setattr(files.settings, "upload_directory", tmp_path)
    monkeypatch.setattr(worker, "Session", database)
    monkeypatch.setattr(harness.settings, "tool_model_allowlist", "")
    buffer = io.BytesIO()
    fact = "행사 장소는 파란 강의실"
    if suffix == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph(fact)
        document.save(buffer)
    elif suffix == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.append([fact, 42])
        book.save(buffer)
    elif suffix == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(6), Inches(1)).text = fact
        deck.save(buffer)
    else:
        buffer.write((json.dumps({"fact": fact}, ensure_ascii=False) if suffix == "json" else
                      f"<fact>{fact}</fact>" if suffix == "xml" else fact).encode())
    client, cid = world
    uploaded = await client.post("/v1/attachments", data={"conversation_id": cid},
                                 files={"file": (f"자료.{suffix}", buffer.getvalue(), "application/octet-stream")})
    assert uploaded.status_code == 200, uploaded.text
    aid = uploaded.json()["id"]
    await worker.process_extraction(aid)
    preview = (await client.get(f"/v1/attachments/{aid}")).json()
    assert preview["status"] == "ready" and fact in preview["preview"]

    async def forbidden(_):
        pytest.fail("Implicit document follow-ups must not send a search")

    monkeypatch.setattr(harness, "search_web", forbidden)
    for index, attachments in enumerate(([aid], [])):
        response = await client.post(f"/v1/conversations/{cid}/runs", headers={"Idempotency-Key": f"office-{index}"},
                                     json={"message": "현재 문서의 행사 장소를 알려줘", "attachment_ids": attachments})
        assert '"status": "completed"' in response.text, response.text
        assert any(fact in m["content"] for m in provider.messages[-1])
    assert len(provider.calls) == 2


async def test_public_config_matches_upload_registry(world):
    client, _ = world
    config = (await client.get("/v1/config")).json()
    assert set(config["attachment_extensions"]) == set(files.ALLOWED)
    assert config["search_default"] == "auto"


async def test_slow_document_is_cancelled_in_time_to_answer_from_ready_sources(world, provider, monkeypatch):
    monkeypatch.setattr(harness.settings, "run_timeout_seconds", 3)
    monkeypatch.setattr(harness.settings, "tool_model_allowlist", "")
    cancelled = asyncio.Event()

    async def page(url):
        if url.endswith("slow.pdf"):
            try:
                await asyncio.sleep(60)
            finally:
                cancelled.set()
        return {"title": url, "url": url, "text": "이미 확보한 본문", "scope": "page"}

    monkeypatch.setattr(harness, "read_url", page)
    client, cid = world
    response = await client.post(f"/v1/conversations/{cid}/runs", json={
        "message": "https://example.com/ready https://example.com/slow.pdf 요약", "search_enabled": False})
    assert '"status": "completed"' in response.text
    assert "RETRIEVAL_TIMEOUT" in response.text and cancelled.is_set()
    assert len(provider.calls) == 1
    assert any("이미 확보한 본문" in message["content"] for message in provider.messages[-1])


async def test_preloaded_documents_do_not_spend_a_planning_call(world, provider, monkeypatch):
    # Even a reviewed planner has no work left after explicit URLs were read.
    async def page(url):
        return {"title": "자료", "url": url, "text": "문서 전체 내용", "scope": "page"}

    monkeypatch.setattr(harness, "read_url", page)
    client, cid = world
    response = await client.post(f"/v1/conversations/{cid}/runs", json={
        "message": "https://example.com/doc 요약해줘", "search_enabled": False})
    assert '"status": "completed"' in response.text and len(provider.calls) == 1


async def test_long_document_fits_selected_text_models_small_context(world, provider, monkeypatch):
    monkeypatch.setattr(harness.settings, "tool_model_allowlist", "")

    async def page(url):
        return {"title": "긴 자료", "url": url, "text": ("일반 안내 내용.\n" * 7000) + "행사장소는 파란강의실입니다.", "scope": "page"}

    monkeypatch.setattr(harness, "read_url", page)
    client, cid = world
    response = await client.post(f"/v1/conversations/{cid}/runs", json={
        "message": "https://example.com/doc 행사장소는 어디?", "search_enabled": False,
        "routing": {"mode": "manual", "provider": "openrouter", "model_id": "test/a"}})
    assert '"status": "completed"' in response.text, response.text
    assert harness.estimate_tokens(provider.messages[-1]) + harness.settings.max_output_tokens <= 8192
    assert any("파란강의실" in m["content"] for m in provider.messages[-1])
    assert "context_truncated" in response.text
