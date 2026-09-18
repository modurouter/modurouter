"""Legacy extraction admission, containment and meaningful conversion contracts."""

import io
import struct
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from modurouter import legacy_office as legacy


def odt_file(tmp_path, *, mime=legacy.ODT, manifest="<manifest/>", content="<document/>"):
    path = tmp_path / "upload"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", mime)
        archive.writestr("META-INF/manifest.xml", manifest)
        archive.writestr("content.xml", content)
    return path


def test_odt_admission_and_type_mismatch(tmp_path):
    legacy.validate_legacy(odt_file(tmp_path), legacy.ODT)
    with pytest.raises(ValueError, match="FILE_TYPE_MISMATCH"):
        legacy.validate_legacy(odt_file(tmp_path, mime="application/zip"), legacy.ODT)


def test_odt_encryption_and_entities_rejected(tmp_path):
    with pytest.raises(ValueError, match="DOCUMENT_ENCRYPTED"):
        legacy.validate_legacy(odt_file(tmp_path, manifest="<manifest><encryption-data/></manifest>"), legacy.ODT)
    with pytest.raises(ValueError, match="DOCUMENT_INVALID"):
        legacy.validate_legacy(odt_file(tmp_path, content='<!DOCTYPE a [<!ENTITY x "payload">]><a>&x;</a>'), legacy.ODT)


def test_odt_duplicate_and_path_traversal_rejected(tmp_path):
    path = odt_file(tmp_path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("../foreign.xml", "<xml/>")
    with pytest.raises(ValueError, match="DOCUMENT_INVALID"):
        legacy.validate_legacy(path, legacy.ODT)
    path = odt_file(tmp_path)
    with pytest.warns(UserWarning), zipfile.ZipFile(path, "a") as archive:
        archive.writestr("content.xml", "<other/>")
    with pytest.raises(ValueError, match="DOCUMENT_INVALID"):
        legacy.validate_legacy(path, legacy.ODT)


def test_odt_expansion_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "MAX_MEMBER", 50)
    with pytest.raises(ValueError, match="DOCUMENT_SIZE_LIMIT"):
        legacy.validate_legacy(odt_file(tmp_path, content="<p>" + "a" * 100 + "</p>"), legacy.ODT)


@pytest.mark.parametrize("mime", [legacy.DOC, legacy.XLS, legacy.PPT, legacy.RTF])
def test_renamed_text_is_rejected(tmp_path, mime):
    path = tmp_path / "upload"
    path.write_text("Not a document", encoding="utf-8")
    with pytest.raises(ValueError, match="FILE_TYPE_MISMATCH"):
        legacy.validate_legacy(path, mime)


def test_rtf_admits_signature(tmp_path):
    path = tmp_path / "upload"
    path.write_bytes(b"{\\rtf1\\ansi Hello}")
    legacy.validate_legacy(path, legacy.RTF)


class FakeOle:
    def __init__(self, streams):
        self.streams = streams

    def __enter__(self):
        return self

    def __exit__(self, *args):
        pass

    def exists(self, name):
        return name in self.streams

    def openstream(self, name):
        return io.BytesIO(self.streams[name])

    def get_size(self, name):
        return len(self.streams[name])


def mock_ole(monkeypatch, streams):
    monkeypatch.setattr(legacy.olefile, "isOleFile", lambda path: True)
    monkeypatch.setattr(legacy.olefile, "OleFileIO", lambda *args, **kwargs: FakeOle(streams))


def test_word_encryption_before_conversion(tmp_path, monkeypatch):
    path = tmp_path / "upload"
    path.write_bytes(b"ole")
    header = bytearray(32)
    struct.pack_into("<HH", header, 0, 0xA5EC, 0x00C1)
    struct.pack_into("<H", header, 10, 0x100)
    mock_ole(monkeypatch, {"WordDocument": bytes(header), "0Table": b""})
    with pytest.raises(ValueError, match="DOCUMENT_ENCRYPTED"):
        legacy.validate_legacy(path, legacy.DOC)
    struct.pack_into("<H", header, 10, 0)
    mock_ole(monkeypatch, {"WordDocument": bytes(header), "0Table": b""})
    legacy.validate_legacy(path, legacy.DOC)


def test_xls_encrypted_workbook(tmp_path, monkeypatch):
    path = tmp_path / "upload"
    path.write_bytes(b"ole")
    data = struct.pack("<HHHH", 0x0809, 0, 0x002F, 0)
    mock_ole(monkeypatch, {"Workbook": data})
    with pytest.raises(ValueError, match="DOCUMENT_ENCRYPTED"):
        legacy.validate_legacy(path, legacy.XLS)


def test_ppt_encryption_token(tmp_path, monkeypatch):
    path = tmp_path / "upload"
    path.write_bytes(b"ole")
    mock_ole(monkeypatch, {"PowerPoint Document": b"", "Current User": b"\0" * 12 + struct.pack("<I", 0xF3D1C4DF)})
    with pytest.raises(ValueError, match="DOCUMENT_ENCRYPTED"):
        legacy.validate_legacy(path, legacy.PPT)


def test_encrypted_ooxml_renamed_legacy_is_rejected(tmp_path, monkeypatch):
    path = tmp_path / "upload"
    path.write_bytes(b"ole")
    mock_ole(monkeypatch, {"EncryptedPackage": b"data"})
    with pytest.raises(ValueError, match="DOCUMENT_ENCRYPTED"):
        legacy.validate_legacy(path, legacy.DOC)


def test_conversion_roundtrip_and_cleanup(tmp_path, monkeypatch):
    from docx import Document

    path = tmp_path / "upload"
    path.write_bytes(b"{\\rtf1\\ansi test}")
    temporary_paths = []

    def convert(source, output, filter_name, directory):
        assert source.suffix == ".rtf"
        assert filter_name == "Office Open XML Text"
        assert source.stat().st_mode & 0o777 == 0o600
        temporary_paths.append(directory)
        document = Document()
        document.add_paragraph("첫 문단: 개시일 9월 18일")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "예산"
        table.cell(0, 1).text = "320만원"
        document.add_paragraph("끝 문단: 담당자 홍길동")
        document.save(output)

    monkeypatch.setattr(legacy, "_convert", convert)
    text = legacy.legacy_text(path, legacy.RTF)
    assert all(term in text for term in ["개시일", "9월 18일", "예산", "320만원", "홍길동", "변환"])
    assert text.index("개시일") < text.index("예산") < text.index("홍길동")
    assert not temporary_paths[0].exists()


def fake_converter(monkeypatch, tmp_path, *, timeout=False, restart=False):
    calls, processes = [], []
    monkeypatch.setattr(legacy.shutil, "which", lambda name: "/usr/bin/soffice")
    monkeypatch.setattr(legacy, "_network_guard", lambda: (lambda: None))
    monkeypatch.setattr(legacy.time, "sleep", lambda seconds: None)
    clock = iter([0, 41] if timeout else [0, 1, 2, 3, 4])
    monkeypatch.setattr(legacy.time, "monotonic", lambda: next(clock))

    class Process:
        def __init__(self, args, **kwargs):
            calls.append((args, kwargs))
            self.terminated = False
            self.helper = "-c" in args
            self.returncode = None if timeout or restart or not self.helper else 0
            self.polls = 0
            if self.helper and not timeout:
                Path(args[-2]).write_bytes(b"output")

        def poll(self):
            self.polls += 1
            if self.helper:
                if restart and self.polls > 1:
                    self.returncode = 0
                return self.returncode
            return 81 if restart and self is processes[0] else None

        def terminate(self):
            self.terminated = True

        def wait(self, timeout):
            return 0

    def popen(args, **kwargs):
        process = Process(args, **kwargs)
        processes.append(process)
        return process

    monkeypatch.setattr(legacy.subprocess, "Popen", popen)
    return calls, processes


def test_converter_uses_private_profile_clean_env_and_shutdown(tmp_path, monkeypatch):
    calls, processes = fake_converter(monkeypatch, tmp_path)
    monkeypatch.setenv("OPENAI_API_KEY", "sentinel-secret")
    legacy._convert(tmp_path / "input.rtf", tmp_path / "output.docx", "Office Open XML Text", tmp_path)
    assert all(process.terminated for process in processes)
    assert len(calls) == 2
    for _, kwargs in calls:
        assert "OPENAI_API_KEY" not in kwargs["env"]
        assert kwargs["env"]["HOME"] == str(tmp_path)
        assert kwargs["close_fds"]
        assert callable(kwargs["preexec_fn"])
    profile = (tmp_path / "profile/user/registrymodifications.xcu").read_text()
    assert 'DisableMacrosExecution' in profile and 'DisableActiveContent' in profile
    assert 'prop("MacroExecutionMode", 0)' in calls[1][0][2]
    assert 'prop("UpdateDocMode", 0)' in calls[1][0][2]
    assert legacy.CONVERSION_TIMEOUT < 60


def test_converter_timeout_always_terminates_office(tmp_path, monkeypatch):
    _, processes = fake_converter(monkeypatch, tmp_path, timeout=True)
    with pytest.raises(ValueError, match="LEGACY_CONVERSION_TIMEOUT"):
        legacy._convert(tmp_path / "input.rtf", tmp_path / "output.docx", "Office Open XML Text", tmp_path)
    assert all(process.terminated for process in processes)


def test_converter_restarts_initial_profile_bootstrap(tmp_path, monkeypatch):
    calls, processes = fake_converter(monkeypatch, tmp_path, restart=True)
    legacy._convert(tmp_path / "input.rtf", tmp_path / "output.docx", "Office Open XML Text", tmp_path)
    assert len(calls) == 3
    assert calls[0][0] == calls[2][0]
    assert processes[1].terminated and processes[2].terminated


def test_missing_converter_actionable(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy.shutil, "which", lambda name: None)
    with pytest.raises(ValueError, match="LEGACY_CONVERTER_UNAVAILABLE"):
        legacy._convert(tmp_path / "input", tmp_path / "output", "filter", tmp_path)


@pytest.mark.skipif(sys.platform != "linux", reason="Production seccomp runs on Linux")
def test_seccomp_denies_network_but_allows_private_unix():
    code = '''
import socket
for family in (socket.AF_INET, socket.AF_INET6):
    try:
        socket.socket(family)
    except PermissionError:
        pass
    else:
        raise AssertionError("network permitted")
socket.socket(socket.AF_UNIX).close()
'''
    subprocess.run([sys.executable, "-c", code], check=True, preexec_fn=legacy._network_guard())

