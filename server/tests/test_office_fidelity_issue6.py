import zipfile
from datetime import datetime

import pytest
from modurouter.office_fidelity import docx_parts, xlsx_parts

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _archive(path, parts):
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


def _doc(path, body, extras=None):
    return _archive(path, {
        "word/document.xml": f'<w:document xmlns:w="{W}" xmlns:r="{R}" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:body>{body}</w:body></w:document>',
        **(extras or {}),
    })


def test_docx_order_nested_tables_textboxes_notes_and_headers(tmp_path):
    path = _doc(tmp_path / "upload", '''
      <w:p><w:r><w:t>첫 문단</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>
      <w:tbl><w:tr><w:tc><w:p><w:r><w:t>항목</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>금액</w:t></w:r></w:p></w:tc></w:tr>
      <w:tr><w:tc><w:p><w:r><w:t>중간 상품</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>300원</w:t></w:r></w:p>
      <w:tbl><w:tr><w:tc><w:p><w:r><w:t>내부표 사실</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
      </w:tc></w:tr></w:tbl>
      <w:p><w:r><w:t>문장 시작</w:t><w:drawing><w:txbxContent><w:p><w:r><w:t>상자 사실</w:t></w:r></w:p></w:txbxContent></w:drawing><w:t>문장 끝</w:t></w:r></w:p>
      <w:p><m:oMath><m:r><m:t>x=42</m:t></m:r></m:oMath></w:p>
      <w:p><w:r><w:t>마지막 문단</w:t><w:endnoteReference w:id="2"/></w:r></w:p>
      <w:sectPr><w:headerReference r:id="rH"/><w:headerReference r:id="rH"/><w:footerReference r:id="rF"/></w:sectPr>''', {
        "word/_rels/document.xml.rels": '<Relationships><Relationship Id="rH" Target="header1.xml"/><Relationship Id="rF" Target="footer1.xml"/></Relationships>',
        "word/header1.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>기밀 머리말</w:t></w:r></w:p></w:hdr>',
        "word/footer1.xml": f'<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>꼬리말 사실</w:t></w:r></w:p></w:ftr>',
        "word/footnotes.xml": f'<w:footnotes xmlns:w="{W}"><w:footnote w:id="-1" w:type="separator"><w:p><w:r><w:t>제외</w:t></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>각주 사실</w:t></w:r></w:p></w:footnote></w:footnotes>',
        "word/endnotes.xml": f'<w:endnotes xmlns:w="{W}"><w:endnote w:id="2"><w:p><w:r><w:t>미주 사실</w:t></w:r></w:p></w:endnote></w:endnotes>',
    })
    text = "\n".join(docx_parts(path))
    assert text.index("첫 문단") < text.index("중간 상품") < text.index("마지막 문단")
    assert "[표 행 1] [열 1] 항목 | [열 2] 금액" in text
    for fact in ("내부표 사실", "상자 사실", "각주 사실", "미주 사실", "기밀 머리말", "꼬리말 사실"):
        assert text.count(fact) == 1
    assert "x=42" in text and "[각주 1]" in text and "[미주 2]" in text
    assert "제외" not in text
    assert "이미지와 차트의 시각 정보" in text


def test_docx_alternate_content_and_deleted_text_not_duplicated(tmp_path):
    path = _doc(tmp_path / "upload", '''<w:p><w:r><w:t>본문</w:t></w:r>
      <mc:AlternateContent><mc:Choice Requires="w"><w:r><w:t>선택 내용</w:t></w:r></mc:Choice>
      <mc:Fallback><w:r><w:t>중복 대체 내용</w:t></w:r></mc:Fallback></mc:AlternateContent>
      <w:del><w:r><w:delText>삭제 내용</w:delText></w:r></w:del>
      </w:p>''')
    text = "\n".join(docx_parts(path))
    assert "선택 내용" in text and "대체 내용" not in text and "삭제 내용" not in text


def test_docx_merged_cells_retain_grid_coordinates(tmp_path):
    path = _doc(tmp_path / "upload", '''<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>병합 제목</w:t></w:r></w:p></w:tc>
      <w:tc><w:p><w:r><w:t>세 번째 열</w:t></w:r></w:p></w:tc></w:tr></w:tbl>''')
    text = "\n".join(docx_parts(path))
    assert "[열 1~2] 병합 제목 | [열 3] 세 번째 열" in text


def test_docx_external_header_is_not_read(tmp_path):
    path = _doc(tmp_path / "upload", '<w:p><w:r><w:t>안전한 문단</w:t></w:r></w:p><w:sectPr><w:headerReference r:id="external"/></w:sectPr>', {
        "word/_rels/document.xml.rels": '<Relationships><Relationship Id="external" TargetMode="External" Target="file:///etc/passwd"/></Relationships>',
    })
    assert "안전한 문단" in "\n".join(docx_parts(path))


def test_docx_xml_depth_bounded(tmp_path):
    path = _doc(tmp_path / "upload", "<w:sdt>" * 130 + "</w:sdt>" * 130)
    with pytest.raises(ValueError, match="OFFICE_SIZE_LIMIT"):
        list(docx_parts(path))


def test_blank_docx_and_xlsx_have_no_notice_only_success(tmp_path):
    from openpyxl import Workbook

    assert list(docx_parts(_doc(tmp_path / "blank.docx", "<w:p/>"))) == []
    workbook = Workbook()
    path = tmp_path / "blank.xlsx"
    workbook.save(path)
    assert list(xlsx_parts(path)) == []


def test_xlsx_sheet_coordinates_dates_units_formulas_and_saved_results(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "요약"
    sheet.append(["날짜", "매출", "비율", "결과", "캐시 없음"])
    sheet.append([datetime(2026, 9, 18), 1234.5, 0.125, "=B2*2", "=1+2"])
    sheet["A2"].number_format = "yyyy-mm-dd"
    sheet["B2"].number_format = '#,##0.0"원"'
    sheet["C2"].number_format = "0.0%"
    sheet["D2"].number_format = '#,##0"원"'
    later = workbook.create_sheet("다른 시트")
    later["C4"] = "후반부 사실"
    path = tmp_path / "upload"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b"<f>B2*2</f><v></v>", b"<f>B2*2</f><v>2469</v>")
    _archive(path, parts)
    text = "\n".join(xlsx_parts(path))
    assert "[시트: 요약]" in text and "[시트: 다른 시트]" in text
    assert "A2: 2026-09-18T00:00:00" in text
    assert 'B2: 1234.5; 표시 형식: #,##0.0"원"; 표시 값: 1,234.5원' in text
    assert "C2: 0.125; 표시 형식: 0.0%; 표시 값: 12.5%" in text
    assert "D2: =B2*2; 수식의 저장된 결과: 2469" in text
    assert "E2: =1+2; 수식의 저장된 결과: [없음: 재계산 필요]" in text
    assert "[행 4] C4: 후반부 사실" in text
    assert "재계산하지 않습니다" in text


def test_xlsx_sparse_extreme_dimensions_are_bounded(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active["A100001"] = "폭탄"
    path = tmp_path / "upload"
    workbook.save(path)
    with pytest.raises(ValueError, match="OFFICE_SIZE_LIMIT"):
        list(xlsx_parts(path))


def test_plain_docx_has_no_unnecessary_visual_notice(tmp_path):
    path = _doc(tmp_path / "upload", '<w:p><w:r><w:t>문단</w:t></w:r></w:p>')
    assert list(docx_parts(path)) == ["문단"]


def test_native_libreoffice_saved_docx_includes_every_document_region():
    from pathlib import Path

    path = Path(__file__).parent / "fixtures/documents/korean-report.docx"
    text = "\n".join(docx_parts(path))
    for fact in ("가람101", "나래202", "다솜303", "라온404", "마루505", "바람606", "1250000", "청록팀"):
        assert text.count(fact) == 1
    assert "[표 행 2] [열 1] 달빛사업 | [열 2] 1250000 | [열 3] 김하늘" in text


def test_native_libreoffice_saved_xlsx_retains_cached_formulas_and_later_sheet():
    from pathlib import Path

    path = Path(__file__).parent / "fixtures/documents/korean-budget.xlsx"
    text = "\n".join(xlsx_parts(path))
    for fact in ("가람101", "나래202", "라온404", "사과707", "아람808", "자람909"):
        assert fact in text
    assert "D2: =B2*C2; 수식의 저장된 결과: 3600" in text
    assert "표시 값: 3,600원" in text
    assert "[시트: 배송현황]" in text and "2026-10-02" in text


def test_zero_padded_codes_are_not_rendered_as_unpadded_values():
    from modurouter.office_fidelity import _display

    assert _display(42, "0000") is None
